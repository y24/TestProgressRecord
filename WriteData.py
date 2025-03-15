from libs import OpenpyxlWrapper as Excel
from openpyxl.worksheet.table import Table, TableStyleInfo
from unicodedata import east_asian_width
from datetime import datetime
from libs import AppConfig

# 列幅調整用の情報
width_dict = {
  'F': 2,   # Fullwidth
  'H': 1,   # Halfwidth
  'W': 2,   # Wide
  'Na': 1,  # Narrow
  'A': 2,   # Ambiguous
  'N': 1    # Neutral
}
Font_depend = 1.2
base_header = ["フォルダ", "ファイル名", "環境名", "日付"]

def convert_to_2d_array(data):
    results = settings["common"]["results"] + ["Completed"]
    header = base_header + results
    out_arr = [header]
    for entry in data:
        file_name = entry.get("file", "")
        if entry["by_env"]:
            for env, env_data in entry.get("by_env", {}).items():
                for date, values in env_data.items():
                    out_arr.append([entry["relative_path"], file_name, env, date] + [values.get(v, 0) for v in results])
        else:
            # 環境別データがない場合は合計データを使用して、環境名は空で出力
            for date, values in entry.get("total_daily", {}).items():
                out_arr.append([entry["relative_path"], file_name, "", date] + [values.get(v, 0) for v in results])
    return out_arr

def create_datatable(ws, sheet_name:str, data, date_labels=["日付"]):
    # ヘッダーを取得
    headers = data[0] if data else []
    date_columns = [idx + 1 for idx, header in enumerate(headers) if header in date_labels]

    # シートのA1からデータ書き込み
    for row_idx, row_data in enumerate(data, start=1):
        for col_idx, value in enumerate(row_data, start=1):
            if col_idx in date_columns and row_idx > 1:  # ヘッダー行を除く
                try:
                    value = datetime.strptime(str(value), "%Y-%m-%d").date()
                except ValueError:
                    pass  # 変換できない場合はそのまま
            ws.cell(row=row_idx, column=col_idx, value=value)

    # テーブル範囲の設定
    max_row = len(data)
    max_col = len(data[0]) if data else 0
    table_ref = f"A1:{chr(64 + max_col)}{max_row}"  # 例: A1:D10

    # テーブル作成
    table = Table(displayName=sheet_name, ref=table_ref)
    style = TableStyleInfo(name="TableStyleLight9", showFirstColumn=False,
                           showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    table.tableStyleInfo = style
    ws.add_table(table)

# ヘッダ名に合致する列の幅を調整
def adjust_colwidth_by_headername(sheet, target_headers:list[str], header_row:int=1):
    for col in sheet.columns:
        # ヘッダ行の値（header_row=1：1行目をヘッダとする）
        header = col[header_row-1].value
        # ヘッダがtarget_headersに合致する列を調整
        if header in target_headers:
            max_length= 1
            max_diameter = 1
            if len(col) > 1:
                column= col[1].column_letter
            else:
                return
            for cell in col:
                diameter = (cell.font.size*Font_depend)/10
                if diameter > max_diameter:
                    max_diameter = diameter
                try:
                    if(cell.value == None) : continue
                    chars = [char for char in str(cell.value)]
                    east_asian_width_list = [east_asian_width(char) for char in chars]
                    width_list = [width_dict[east_asian_width] for east_asian_width in east_asian_width_list]
                    if sum(width_list) > max_length:
                        max_length= sum(width_list)
                except:
                    pass
                sheet.column_dimensions[column].width= max_length*max_diameter*0.9

def execute(data, file_path, sheet_name):
    global settings

    # 設定読み込み
    settings = AppConfig.load_settings()

    # データを書込用に変換
    converted_data = convert_to_2d_array(data)

    # ブック読み込み
    wb = Excel.load(file_path=file_path)
    
    # 既存のデータシートを削除して新規作成
    ws = Excel.create_sheet(workbook=wb, sheet_name=sheet_name, overwrite=True)

    # データがあれば書き込み
    if len(converted_data) > 1:        
        # データを書き込んでテーブルを作成
        create_datatable(ws, sheet_name, converted_data)
        # 列幅の調整
        adjust_colwidth_by_headername(sheet=ws, target_headers=base_header, header_row=1)
    else:
        raise ValueError("書き込み可能なデータが1件もありません。")

    # 保存
    try:
        wb.save(file_path)
        return True
    except PermissionError:
        raise PermissionError("書込先のファイルに権限がないか、読み取り専用の可能性があります。\nファイルを開いている場合は閉じてやり直してください。")
